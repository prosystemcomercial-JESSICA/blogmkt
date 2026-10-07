# Gera a planilha "Ficha técnica de padaria" oferecida no blog em troca do cadastro.
# Uso (a partir da pasta blog/):  python _gerador/ficha_tecnica_xlsx.py
import pathlib

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BLOG = pathlib.Path(__file__).resolve().parent.parent
SAIDA = BLOG / "assets" / "downloads" / "ficha-tecnica-padaria-prosystem.xlsx"
LOGO = BLOG / "assets" / "logo" / "logo-h.png"

AZUL = "0A2E87"
AZUL_SITE = "4379BC"
NOITE = "081330"
CINZA = "667085"
LINHA = "D0D5DD"
PREENCHER = "EAF2FB"   # células para o cliente preencher (azul-claro; a marca não usa amarelo)
RESULTADO = "F2F4F7"

F = "Arial"
fonte = lambda **k: Font(name=F, **k)
borda = Border(*(Side(style="thin", color=LINHA),) * 4)
preencher = PatternFill("solid", fgColor=PREENCHER)
resultado = PatternFill("solid", fgColor=RESULTADO)
cabecalho = PatternFill("solid", fgColor=AZUL)

MOEDA = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
NUM3 = '#,##0.000;-#,##0.000;"-"'
INT = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'

LINHAS_ING = 14           # linhas de ingredientes
L0 = 15                   # primeira linha de ingredientes

EXEMPLO = {
    "produto": "Pão francês",
    "categoria": "Pães",
    "responsavel": "Padeiro do turno da manhã",
    "ingredientes": [
        ("Farinha de trigo", 10.0, "kg", 5.50),
        ("Água", 6.0, "kg", 0.0),
        ("Fermento biológico fresco", 0.3, "kg", 30.00),
        ("Sal", 0.2, "kg", 3.00),
        ("Açúcar", 0.1, "kg", 5.00),
        ("Melhorador", 0.1, "kg", 60.00),
    ],
    "peso_peca": 60,
    "real": 262,
    "embalagem": 0.0,
    "outros_pct": 0.25,
    "mao_obra": 45.00,
    "impostos": 0.06,
    "margem": 0.35,
    "preco_atual": 0.80,
}


def celula(ws, ref, valor=None, *, negrito=False, cor="101828", tam=10, fmt=None, fill=None,
           alinh=None, b=False, italico=False):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.font = fonte(bold=negrito, color=cor, size=tam, italic=italico)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if alinh:
        c.alignment = alinh
    if b:
        c.border = borda
    return c


def montar(ws, ex):
    ws.sheet_view.showGridLines = False
    larguras = {"A": 2, "B": 34, "C": 14, "D": 12, "E": 18, "F": 18, "G": 2, "H": 30}
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w

    # Cabeçalho
    if LOGO.exists():
        img = XLImage(str(LOGO))
        img.width, img.height = 210, 30
        ws.add_image(img, "B2")
    ws.row_dimensions[2].height = 26
    celula(ws, "B4", "Ficha técnica de produção", negrito=True, cor=NOITE, tam=16)
    celula(ws, "B5", "Padronize a receita, confira o rendimento de cada fornada e descubra o custo e o preço de cada unidade.",
           cor=CINZA, tam=10)
    celula(ws, "B6", "Preencha apenas as células em azul-claro. As demais são calculadas automaticamente.",
           cor=AZUL, tam=9, italico=True)
    celula(ws, "E6", "", fill=preencher, b=True)
    celula(ws, "F6", "célula para preencher", cor=CINZA, tam=9)

    # Dados do produto
    for i, (rot, chave) in enumerate([("Produto", "produto"), ("Categoria", "categoria"),
                                       ("Responsável pela receita", "responsavel")]):
        r = 8 + i
        celula(ws, f"B{r}", rot, negrito=True, cor=NOITE)
        ws.merge_cells(f"C{r}:F{r}")
        celula(ws, f"C{r}", ex.get(chave, "") if ex else None, fill=preencher, b=True)
        for col in "DEF":
            ws[f"{col}{r}"].border = borda
            ws[f"{col}{r}"].fill = preencher
    celula(ws, "B11", "Data da última revisão", negrito=True, cor=NOITE)
    celula(ws, "C11", "=TODAY()" if ex else None, fmt="dd/mm/yyyy", fill=preencher, b=True)

    # Tabela de ingredientes
    celula(ws, "B13", "1. Ingredientes da receita", negrito=True, cor=AZUL, tam=12)
    cab = ["Ingrediente", "Quantidade", "Unidade", "Preço por unidade", "Custo na receita"]
    for j, t in enumerate(cab):
        c = celula(ws, f"{'BCDEF'[j]}14", t, negrito=True, cor="FFFFFF", fill=cabecalho, b=True,
                   alinh=Alignment(horizontal="center", vertical="center", wrap_text=True))
    ws.row_dimensions[14].height = 30
    ws["E14"].comment = Comment("Preço pago por kg, litro ou unidade, conforme a coluna Unidade. Use o valor da última nota fiscal de compra.", "ProSystem")

    dv = DataValidation(type="list", formula1='"kg,L,un"', allow_blank=True)
    ws.add_data_validation(dv)
    for k in range(LINHAS_ING):
        r = L0 + k
        dados = ex["ingredientes"][k] if ex and k < len(ex["ingredientes"]) else None
        celula(ws, f"B{r}", dados[0] if dados else None, fill=preencher, b=True)
        celula(ws, f"C{r}", dados[1] if dados else None, fmt=NUM3, fill=preencher, b=True)
        celula(ws, f"D{r}", dados[2] if dados else None, fill=preencher, b=True, alinh=Alignment(horizontal="center"))
        dv.add(f"D{r}")
        celula(ws, f"E{r}", dados[3] if dados else None, fmt=MOEDA, fill=preencher, b=True)
        celula(ws, f"F{r}", f"=IF(OR(C{r}=\"\",E{r}=\"\"),0,C{r}*E{r})", fmt=MOEDA, fill=resultado, b=True)
    LF = L0 + LINHAS_ING - 1
    rt = LF + 1
    celula(ws, f"B{rt}", "Custo total dos ingredientes", negrito=True, cor=NOITE, b=True)
    celula(ws, f"C{rt}", f'=SUMIF(D{L0}:D{LF},"kg",C{L0}:C{LF})+SUMIF(D{L0}:D{LF},"L",C{L0}:C{LF})',
           fmt=NUM3, fill=resultado, b=True, negrito=True)
    celula(ws, f"D{rt}", "kg de massa", cor=CINZA, tam=9, b=True)
    celula(ws, f"E{rt}", "", b=True)
    celula(ws, f"F{rt}", f"=SUM(F{L0}:F{LF})", fmt=MOEDA, fill=resultado, b=True, negrito=True)
    ws[f"C{rt}"].comment = Comment("Soma dos ingredientes em kg e em litros (considera 1 L de água ≈ 1 kg). Itens em 'un' não entram na massa.", "ProSystem")

    # Rendimento
    r0 = rt + 2
    celula(ws, f"B{r0}", "2. Rendimento da fornada", negrito=True, cor=AZUL, tam=12)
    linhas = [
        ("Peso da peça crua (g)", ex["peso_peca"] if ex else None, INT, True, None),
        ("Rendimento previsto (unidades)", f'=IF(C{r0+1}>0,INT(C{rt}*1000/C{r0+1}),0)', INT, False,
         "Massa total dividida pelo peso da peça crua."),
        ("Rendimento real da última fornada (unidades)", ex["real"] if ex else None, INT, True,
         "Conte quantas unidades realmente saíram. Diferenças frequentes indicam problema de pesagem, divisora ou forno."),
        ("Diferença entre real e previsto", f'=IF(AND(C{r0+2}>0,C{r0+3}>0),(C{r0+3}-C{r0+2})/C{r0+2},0)', PCT, False,
         "Negativo = rendeu menos do que a receita prevê."),
    ]
    for i, (rot, val, fmt, entrada, nota) in enumerate(linhas):
        r = r0 + 1 + i
        celula(ws, f"B{r}", rot, cor=NOITE, b=True)
        celula(ws, f"C{r}", val, fmt=fmt, fill=preencher if entrada else resultado, b=True, negrito=not entrada)
        if nota:
            ws[f"C{r}"].comment = Comment(nota, "ProSystem")
    unid = f"IF(C{r0+3}>0,C{r0+3},C{r0+2})"   # usa o real; se vazio, o previsto

    # Custo por unidade
    c0 = r0 + 6
    celula(ws, f"B{c0}", "3. Custo por unidade", negrito=True, cor=AZUL, tam=12)
    custos = [
        ("Ingredientes por unidade", f"=IF({unid}>0,F{rt}/{unid},0)", MOEDA, False,
         "Custo dos ingredientes dividido pelas unidades (usa o rendimento real; se vazio, o previsto)."),
        ("Embalagem por unidade (R$)", ex["embalagem"] if ex else None, MOEDA, True, "Saquinho, etiqueta, bandeja. Deixe 0 se não houver."),
        ("Energia, gás e outros custos (% sobre ingredientes)", ex["outros_pct"] if ex else None, PCT, True,
         "Estimativa do peso de energia, gás e itens indiretos. Ajuste com os seus números."),
        ("Mão de obra da fornada (R$)", ex["mao_obra"] if ex else None, MOEDA, True,
         "Custo das horas da equipe dedicadas a esta fornada."),
        ("Custo total por unidade", None, MOEDA, False, None),
    ]
    for i, (rot, val, fmt, entrada, nota) in enumerate(custos):
        r = c0 + 1 + i
        celula(ws, f"B{r}", rot, cor=NOITE, b=True, negrito=(i == 4))
        celula(ws, f"C{r}", val, fmt=fmt, fill=preencher if entrada else resultado, b=True, negrito=not entrada)
        if nota:
            ws[f"C{r}"].comment = Comment(nota, "ProSystem")
    ct = c0 + 5
    ws[f"C{ct}"] = f"=C{c0+1}+C{c0+2}+C{c0+1}*C{c0+3}+IF({unid}>0,C{c0+4}/{unid},0)"

    # Preço
    p0 = c0 + 7
    celula(ws, f"B{p0}", "4. Preço de venda", negrito=True, cor=AZUL, tam=12)
    precos = [
        ("Impostos sobre a venda (% do preço)", ex["impostos"] if ex else None, PCT, True,
         "Alíquota efetiva dos impostos sobre a venda. Confirme com o seu contador."),
        ("Margem desejada (% do preço)", ex["margem"] if ex else None, PCT, True,
         "Quanto do preço de venda deve sobrar depois de custos e impostos."),
        ("Preço sugerido por unidade", f'=IF((1-C{p0+1}-C{p0+2})>0,C{ct}/(1-C{p0+1}-C{p0+2}),0)', MOEDA, False,
         "Custo total dividido por (1 - impostos - margem)."),
        ("Preço que você pratica hoje (opcional)", ex["preco_atual"] if ex else None, MOEDA, True, None),
        ("Margem real com o preço de hoje", f'=IF(C{p0+4}>0,(C{p0+4}*(1-C{p0+1})-C{ct})/C{p0+4},0)', PCT, False,
         "O que sobra do preço atual depois de custos e impostos."),
    ]
    for i, (rot, val, fmt, entrada, nota) in enumerate(precos):
        r = p0 + 1 + i
        celula(ws, f"B{r}", rot, cor=NOITE, b=True, negrito=(i in (2, 4)))
        celula(ws, f"C{r}", val, fmt=fmt, fill=preencher if entrada else resultado, b=True, negrito=not entrada)
        if nota:
            ws[f"C{r}"].comment = Comment(nota, "ProSystem")

    # Caixa lateral
    celula(ws, "H8", "Como usar", negrito=True, cor=NOITE, tam=11)
    passos = [
        "1. Pese tudo numa produção real e anote em kg ou L.",
        "2. Use o preço da última nota fiscal de compra.",
        "3. Conte as unidades que realmente saíram da fornada.",
        "4. Compare o rendimento real com o previsto.",
        "5. Revise a ficha sempre que um insumo mudar de preço.",
    ]
    for i, t in enumerate(passos):
        celula(ws, f"H{9 + i}", t, cor="344054", alinh=Alignment(wrap_text=True, vertical="top"))
    celula(ws, "H16", "Cansou de atualizar preço à mão?", negrito=True, cor=AZUL, tam=11)
    celula(ws, "H17", "No sistema ProSystem, a entrada da nota fiscal pelo XML atualiza o custo dos insumos e a formação de preço acompanha.",
           cor="344054", alinh=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[17].height = 54
    c = celula(ws, "H19", "Fale com um especialista: (27) 99752-1370", negrito=True, cor=AZUL_SITE)
    c.hyperlink = "https://wa.me/5527997521370?text=Ol%C3%A1!%20Baixei%20a%20ficha%20t%C3%A9cnica%20do%20blog%20e%20quero%20conhecer%20o%20sistema%20para%20padaria."
    c = celula(ws, "H20", "prosystemnet.com/blog", cor=AZUL_SITE)
    c.hyperlink = "https://prosystemnet.com/blog/ficha-tecnica-padaria-padronizar-receitas-rendimento/"

    if ex:
        celula(ws, f"B{p0 + 7}", "Valores ilustrativos para mostrar o cálculo. Use os números da sua padaria na aba \"Minha ficha\".",
               cor=CINZA, tam=9, italico=True)
    ws.freeze_panes = "A8"
    ws.print_area = f"A1:H{p0 + 7}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


wb = Workbook()
ws1 = wb.active
ws1.title = "Minha ficha"
montar(ws1, None)
ws2 = wb.create_sheet("Exemplo - Pão francês")
montar(ws2, EXEMPLO)
wb.active = 0
wb.properties.creator = "ProSystem Sistemas"
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.properties.title = "Ficha técnica de produção para padaria"
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print(SAIDA)
